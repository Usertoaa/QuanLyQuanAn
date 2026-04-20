"""
Module xử lý import dữ liệu các món ăn từ file Excel
"""
import pandas as pd
from openpyxl import load_workbook
from django.core.files.uploadedfile import UploadedFile
from django.db import transaction
from .models import Dish, Restaurant


class DishImportHandler:
    """Xử lý import dữ liệu các món ăn từ file Excel"""
    
    # Các cột bắt buộc trong file Excel
    REQUIRED_COLUMNS = ['Tên Món', 'Giá Tiền']
    
    # Các cột tùy chọn
    OPTIONAL_COLUMNS = ['Mô Tả', 'Còn Món', 'Dùng Làm Giá Đại Diện']
    
    def __init__(self, excel_file: UploadedFile, restaurant: Restaurant):
        """
        Khởi tạo handler
        
        Args:
            excel_file: UploadedFile từ form
            restaurant: Instance của Restaurant model
        """
        self.excel_file = excel_file
        self.restaurant = restaurant
        self.errors = []
        self.warnings = []
        self.success_count = 0
        self.failed_count = 0
        self.data = None
    
    def validate_file(self) -> bool:
        """Kiểm tra tính hợp lệ của file Excel"""
        try:
            # Đọc file Excel
            self.data = pd.read_excel(self.excel_file, sheet_name=0)
            
            # Kiểm tra file có dữ liệu không
            if self.data.empty:
                self.errors.append("File Excel không có dữ liệu")
                return False
            
            # Kiểm tra các cột bắt buộc
            missing_columns = []
            for col in self.REQUIRED_COLUMNS:
                if col not in self.data.columns:
                    missing_columns.append(col)
            
            if missing_columns:
                self.errors.append(
                    f"Thiếu các cột bắt buộc: {', '.join(missing_columns)}"
                )
                return False
            
            return True
            
        except Exception as e:
            self.errors.append(f"Lỗi khi đọc file Excel: {str(e)}")
            return False
    
    def parse_price(self, price_value) -> int:
        """
        Chuyển đổi giá tiền từ các định dạng khác nhau
        
        Args:
            price_value: Giá tiền có thể là int, float, hoặc string
            
        Returns:
            Giá tiền dưới dạng integer, hoặc None nếu không hợp lệ
        """
        try:
            if pd.isna(price_value):
                return None
            
            # Xử lý nếu là string
            if isinstance(price_value, str):
                # Loại bỏ khoảng trắng và ký tự đặc biệt
                price_str = price_value.strip().replace('.', '').replace(',', '')
                price_value = int(float(price_str))
            else:
                price_value = int(float(price_value))
            
            if price_value < 0:
                return None
            
            return price_value
            
        except (ValueError, TypeError):
            return None
    
    def parse_boolean(self, value, default=True) -> bool:
        """
        Chuyển đổi giá trị thành boolean
        
        Args:
            value: Giá trị cần chuyển (string, bool, int, etc.)
            default: Giá trị mặc định nếu không xác định
            
        Returns:
            Boolean value
        """
        if pd.isna(value):
            return default
        
        if isinstance(value, bool):
            return value
        
        if isinstance(value, (int, float)):
            return bool(value)
        
        if isinstance(value, str):
            true_values = ['yes', 'y', 'true', '1', 'có', 'c', 'đúng']
            return value.strip().lower() in true_values
        
        return default
    
    def import_dishes(self) -> dict:
        """
        Thực hiện import dữ liệu các món ăn
        
        Returns:
            Dictionary chứa thông tin kết quả import
        """
        if not self.validate_file():
            return self._get_result()
        
        try:
            with transaction.atomic():
                # Lặp qua từng dòng dữ liệu
                for index, row in self.data.iterrows():
                    try:
                        self._import_single_dish(row, index + 2)  # +2 vì row 0 là header, index bắt đầu từ 0
                    except Exception as e:
                        self.errors.append(f"Lỗi ở hàng {index + 2}: {str(e)}")
                        self.failed_count += 1
        
        except Exception as e:
            self.errors.append(f"Lỗi khi lưu dữ liệu: {str(e)}")
        
        return self._get_result()
    
    def _import_single_dish(self, row, row_number: int):
        """
        Import một mon ăn từ dòng dữ liệu
        
        Args:
            row: Một dòng từ DataFrame
            row_number: Số dòng (để báo lỗi)
        """
        # Lấy tên món (bắt buộc)
        dish_name = str(row['Tên Món']).strip()
        if not dish_name or pd.isna(row['Tên Món']):
            raise ValueError("Tên món không được để trống")
        
        # Lấy giá tiền (bắt buộc)
        price = self.parse_price(row['Giá Tiền'])
        if price is None:
            raise ValueError("Giá tiền không hợp lệ")
        
        # Lấy mô tả (tùy chọn)
        description = ""
        if 'Mô Tả' in row and not pd.isna(row['Mô Tả']):
            description = str(row['Mô Tả']).strip()
        
        # Lấy trạng thái còn món (tùy chọn, mặc định là True)
        is_available = True
        if 'Còn Món' in row:
            is_available = self.parse_boolean(row['Còn Món'], default=True)
        
        # Lấy giá đại diện (tùy chọn, mặc định là False)
        is_price_representative = False
        if 'Dùng Làm Giá Đại Diện' in row:
            is_price_representative = self.parse_boolean(row['Dùng Làm Giá Đại Diện'], default=False)
        
        # Tạo hoặc cập nhật Dish
        dish, created = Dish.objects.update_or_create(
            restaurant=self.restaurant,
            name=dish_name,
            defaults={
                'description': description,
                'price': price,
                'is_available': is_available,
                'is_price_representative': is_price_representative,
            }
        )
        
        if created:
            self.success_count += 1
        else:
            self.warnings.append(f"Hàng {row_number}: Món '{dish_name}' đã tồn tại, cập nhật thông tin")
            self.success_count += 1
    
    def _get_result(self) -> dict:
        """Trả về kết quả import"""
        return {
            'success': len(self.errors) == 0,
            'errors': self.errors,
            'warnings': self.warnings,
            'success_count': self.success_count,
            'failed_count': self.failed_count,
            'total': self.success_count + self.failed_count,
            'message': self._get_message()
        }
    
    def _get_message(self) -> str:
        """Tạo thông báo kết quả"""
        if not self.errors and self.success_count > 0:
            msg = f"✅ Import thành công {self.success_count} món ăn"
            if self.warnings:
                msg += f" (có {len(self.warnings)} cảnh báo)"
            return msg
        elif self.errors:
            return f"❌ Import không thành công. Lỗi: {len(self.errors)} lỗi"
        else:
            return "⚠️ Không có dữ liệu được import"


def create_sample_excel_template(output):
    """
    Tạo file Excel mẫu để import dữ liệu
    
    Args:
        output: Có thể là đường dẫn file (str) hoặc BytesIO object
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    
    # Tạo workbook mới
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Thực Đơn"
    
    # Định dạng header
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    center_alignment = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Thêm header
    headers = ['Tên Món', 'Giá Tiền', 'Mô Tả', 'Còn Món', 'Dùng Làm Giá Đại Diện']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_alignment
        cell.border = border
    
    # Thêm dòng hướng dẫn (dòng 2)
    instructions = [
        'Tên món ăn (bắt buộc)',
        'Giá tiền VNĐ (bắt buộc)',
        'Mô tả chi tiết về món',
        'Có (Yes/Y/True/1/Có/C/Đúng) hoặc Không (No/N/False/0)',
        'Có (Yes/Y/True/1/Có/C/Đúng) hoặc Không (No/N/False/0)'
    ]
    
    instruction_font = Font(italic=True, size=9, color="666666")
    instruction_fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
    
    for col_num, instruction in enumerate(instructions, 1):
        cell = ws.cell(row=2, column=col_num)
        cell.value = instruction
        cell.font = instruction_font
        cell.fill = instruction_fill
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = border
    
    # Thêm dữ liệu mẫu (dòng 3-5)
    sample_data = [
        ['Phở Bò', 50000, 'Phở bò nóng với nước dùng thơm ngon', 'Có', 'Có'],
        ['Gà Rán Giòn', 45000, 'Gà rán giòn sốt cà chua chua ngon', 'Có', 'Không'],
        ['Cơm Tấm Sườn Nướng', 35000, 'Cơm tấm, sườn nướng, ăn kèm dưa leo và trứng', 'Có', 'Không'],
    ]
    
    data_font = Font(size=10)
    
    for row_num, row_data in enumerate(sample_data, 3):
        for col_num, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.value = value
            cell.font = data_font
            cell.border = border
            
            # Format cột giá tiền
            if col_num == 2:
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 35
    ws.column_dimensions['D'].width = 18
    ws.column_dimensions['E'].width = 22
    
    # Set row heights
    ws.row_dimensions[1].height = 25
    ws.row_dimensions[2].height = 40
    
    # Thêm sheet hướng dẫn
    ws_guide = wb.create_sheet("Hướng Dẫn")
    
    guide_content = [
        ["HƯỚNG DẪN NHẬP DỮ LIỆU CÁC MÓN ĂN"],
        [""],
        ["1. YÊU CẦU CÁC CỘT:"],
        ["   - Tên Món (bắt buộc): Tên của món ăn"],
        ["   - Giá Tiền (bắt buộc): Giá tiền tính bằng VNĐ"],
        [""],
        ["2. CÁC CỘT TÙY CHỌN:"],
        ["   - Mô Tả: Mô tả chi tiết về món ăn"],
        ["   - Còn Món: Trạng thái còn/hết món (Có/Không)"],
        ["   - Dùng Làm Giá Đại Diện: Sử dụng giá này làm giá đại diện (Có/Không)"],
        [""],
        ["3. ĐỊNH DẠNG DỮ LIỆU:"],
        ["   - Tên Món: Text, không được để trống"],
        ["   - Giá Tiền: Số nguyên dương (vd: 50000, 45000, 35000)"],
        ["   - Mô Tả: Text, có thể để trống"],
        ["   - Còn Món: Chọn từ (Có, Yes, Y, True, 1) hoặc (Không, No, N, False, 0)"],
        ["   - Dùng Làm Giá Đại Diện: Chọn từ (Có, Yes, Y, True, 1) hoặc (Không, No, N, False, 0)"],
        [""],
        ["4. GHI CHÚ:"],
        ["   - Nếu một món đã tồn tại, dữ liệu sẽ được cập nhật"],
        ["   - Tối đa 1000 dòng dữ liệu trong một file"],
        ["   - File phải có định dạng .xlsx hoặc .xls"],
        ["   - Kích thước file không được vượt quá 5MB"],
    ]
    
    for row_num, row_data in enumerate(guide_content, 1):
        for col_num, value in enumerate(row_data, 1):
            cell = ws_guide.cell(row=row_num, column=col_num)
            cell.value = value
            if row_num == 1:
                cell.font = Font(bold=True, size=14)
            elif any(x in str(value) for x in ['YÊU CẦU', 'CÁC CỘT', 'ĐỊNH DẠNG', 'GHI CHÚ']):
                cell.font = Font(bold=True, size=11)
    
    ws_guide.column_dimensions['A'].width = 80
    
    # Lưu workbook
    wb.save(output)
